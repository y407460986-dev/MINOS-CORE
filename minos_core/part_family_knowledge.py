"""Part-family knowledge conversion and loading.

This module turns the three legacy XLSX seed workbooks into compact,
deduplicated knowledge payloads that can be loaded by family key or path.
"""

from __future__ import annotations

import json
import os
import re
from pathlib import Path
from typing import Any

from minos_core._paths import configs_dir

SCHEMA_VERSION = "minos-core.part_family_knowledge.v1"
INDEX_SCHEMA_VERSION = "minos-core.part_family_index.v1"


def _config_dir() -> Path:
    """Return the part-families config directory."""
    return configs_dir() / "knowledge" / "part_families"


def _knowledge_root() -> Path:
    """Return the parent of the part-families config directory."""
    return _config_dir().parent


def _process_inference_dir() -> Path:
    """Return the process-inference knowledge directory."""
    return _knowledge_root() / "process_inference"


def _index_path() -> Path:
    """Return the path to the part-family index JSON."""
    return _config_dir() / "index.json"


CONFIG_DIR = _config_dir()
KNOWLEDGE_ROOT = _knowledge_root()
PROCESS_INFERENCE_DIR = _process_inference_dir()
INDEX_PATH = _index_path()


def _default_seed_dir() -> Path:
    """Return the default seed directory, overridable via env var MINOS_KNOWLEDGE_SEED_DIR."""
    env = os.environ.get("MINOS_KNOWLEDGE_SEED_DIR")
    if env:
        return Path(env)
    return configs_dir() / "knowledge_seed"


def _default_seed_directories() -> dict[str, Path]:
    """Return a mapping of family keys to their default seed directories."""
    seed = _default_seed_dir()
    return {
        "stamping": seed / "冲压件",
        "limiter": seed / "限位器",
        "door_hinge": seed / "门铰链",
    }


_DEFAULT_SEED = _default_seed_dir()
DEFAULT_SEED_DIRECTORIES: dict[str, Path] = _default_seed_directories()

DEFAULT_FAMILY_NAMES: dict[str, str] = {
    "stamping": "冲压件",
    "limiter": "限位器",
    "door_hinge": "门铰链",
    "stamping_bracket": "冲压支架",
    "deep_drawn_part": "拉深件",
    "spot_welded_part": "点焊组件",
    "cold_headed_part": "冷镦件",
    "injection_molded": "注塑件",
    "insert_molded": "嵌件注塑",
    "die_cast_part": "压铸件",
    "rubber_bonded": "橡胶硫化件",
    "welded_assembly": "焊接总成",
    "machined_shaft": "机加工轴类",
    "machined_housing": "机加工壳体",
    "assembly_module": "装配模块",
}

FAMILY_ALIASES: dict[str, str] = {
    "冲压件": "stamping",
    "冲压": "stamping",
    "stamping": "stamping",
    "stamp": "stamping",
    "stamped_component": "stamping",
    "press": "stamping",
    "限位器": "limiter",
    "限位": "limiter",
    "limiter": "limiter",
    "limit": "limiter",
    "door_limiter": "limiter",
    "check_limiter": "limiter",
    "门铰链": "door_hinge",
    "铰链": "door_hinge",
    "door_hinge": "door_hinge",
    "door-hinge": "door_hinge",
    "doorhinge": "door_hinge",
    "hinge": "door_hinge",
    "冲压支架": "stamping_bracket",
    "支架": "stamping_bracket",
    "stamping_bracket": "stamping_bracket",
    "bracket": "stamping_bracket",
    "拉深件": "deep_drawn_part",
    "拉深": "deep_drawn_part",
    "deep_drawn_part": "deep_drawn_part",
    "deep_draw": "deep_drawn_part",
    "点焊组件": "spot_welded_part",
    "点焊": "spot_welded_part",
    "spot_welded_part": "spot_welded_part",
    "spot_weld": "spot_welded_part",
    "冷镦件": "cold_headed_part",
    "冷镦": "cold_headed_part",
    "cold_headed_part": "cold_headed_part",
    "cold_heading": "cold_headed_part",
    "注塑件": "injection_molded",
    "注塑": "injection_molded",
    "injection_molded": "injection_molded",
    "injection_mold": "injection_molded",
    "嵌件注塑": "insert_molded",
    "insert_molded": "insert_molded",
    "insert_mold": "insert_molded",
    "压铸件": "die_cast_part",
    "压铸": "die_cast_part",
    "die_cast_part": "die_cast_part",
    "die_cast": "die_cast_part",
    "橡胶硫化件": "rubber_bonded",
    "橡胶": "rubber_bonded",
    "rubber_bonded": "rubber_bonded",
    "rubber": "rubber_bonded",
    "焊接总成": "welded_assembly",
    "焊接": "welded_assembly",
    "welded_assembly": "welded_assembly",
    "weld_assembly": "welded_assembly",
    "机加工轴类": "machined_shaft",
    "轴类": "machined_shaft",
    "machined_shaft": "machined_shaft",
    "shaft": "machined_shaft",
    "机加工壳体": "machined_housing",
    "壳体": "machined_housing",
    "machined_housing": "machined_housing",
    "housing": "machined_housing",
    "装配模块": "assembly_module",
    "装配": "assembly_module",
    "assembly_module": "assembly_module",
    "assembly": "assembly_module",
}

PROFILES: dict[str, dict[str, Any]] = {
    "door_hinge": {
        "process_flow": {
            "sheets": ["过程流程图"],
            "start": 8,
            "end": 24,
            "step": "A",
            "operation": "A",
            "function": ["F"],
            "special": ["H", "K"],
        },
        "pfmea": {
            "sheets": ["PFMEA"],
            "start": 16,
            "end": 58,
            "step": "A",
            "function": ["A"],
            "failure_mode": "B",
            "effect": "D",
            "severity": "F",
            "classification": "G",
            "cause": "H",
            "occurrence": "I",
            "prevention": "J",
            "detection_control": "K",
            "detection": "L",
            "action_priority": "",
            "inherit": ["A"],
        },
        "control_plan": {
            "sheets": ["控制计划"],
            "start": 16,
            "end": 57,
            "step": "A",
            "operation": "B",
            "characteristic": ["E", "F"],
            "classification": "G",
            "specification": "H",
            "measurement": "I",
            "sample_size": "J",
            "sample_frequency": "K",
            "control_method": "L",
            "reaction_plan": "N",
            "inherit": ["A", "B"],
        },
    },
    "limiter": {
        "process_flow": {
            "sheets": ["4过程流程图", "过程流程图"],
            "start": 10,
            "end": 34,
            "step": "A",
            "operation": "B",
            "function": ["G"],
            "special": ["H", "I", "J", "K"],
        },
        "pfmea": {
            "sheets": ["6PFMEA", "PFMEA"],
            "start": 10,
            "end": 71,
            "step": "B",
            "function": ["D", "E", "F"],
            "failure_mode": "I",
            "effect": "G",
            "severity": "H",
            "classification": "P",
            "cause": "J",
            "occurrence": "L",
            "prevention": "K",
            "detection_control": "M",
            "detection": "N",
            "action_priority": "O",
            "inherit": ["B", "D", "E", "F"],
        },
        "control_plan": {
            "sheets": ["7控制计划-过程", "控制计划"],
            "start": 12,
            "end": 81,
            "step": "A",
            "operation": "B",
            "characteristic": ["I", "L"],
            "classification": "O",
            "specification": "Q",
            "measurement": "U",
            "sample_size": "Y",
            "sample_frequency": "Z",
            "control_method": "AB",
            "reaction_plan": "AE",
            "inherit": ["A", "B"],
        },
    },
    "stamping": {
        "process_flow": {
            "sheets": ["7 过程流程图", "过程流程图"],
            "start": 9,
            "end": 20,
            "step": "A",
            "operation": "A",
            "function": ["A"],
            "special": ["I", "J", "K"],
        },
        "pfmea": {
            "sheets": ["8 PFMEA", "PFMEA"],
            "start": 16,
            "end": 69,
            "step": "A",
            "function": ["B"],
            "failure_mode": "C",
            "effect": "D",
            "severity": "E",
            "classification": "F",
            "cause": "G",
            "occurrence": "H",
            "prevention": "I",
            "detection_control": "J",
            "detection": "K",
            "action_priority": "",
            "inherit": ["A", "B"],
        },
        "control_plan": {
            "sheets": ["9 控制计划", "控制计划"],
            "start": 13,
            "end": 70,
            "step": "A",
            "operation": "B",
            "characteristic": ["F", "G", "H"],
            "classification": "I",
            "specification": "J",
            "measurement": "K",
            "sample_size": "L",
            "sample_frequency": "M",
            "control_method": "N",
            "reaction_plan": "O",
            "inherit": ["A", "B", "C"],
        },
    },
}


def convert_seed_workbook_to_family_knowledge(
    path: str | Path,
    family_key: str,
    family_name: str | None = None,
) -> dict[str, Any]:
    """Convert one XLSX seed workbook into part-family knowledge.

    ``path`` may point directly to a workbook or to a directory containing one
    ``.xlsx``/``.xlsm`` seed workbook. ``openpyxl`` is required and no fallback
    parser is used.

    Args:
        path: Path to a workbook file or a directory containing one seed workbook.
        family_key: Normalised key identifying the part family (e.g. ``stamping``).
        family_name: Optional human-readable family name; defaults from
            ``DEFAULT_FAMILY_NAMES``.

    Returns:
        A knowledge payload dict with schema_version, family_key, family_name,
        process_route_templates, pfmea_patterns, control_plan_patterns,
        special_characteristic_patterns, and source_refs.

    Raises:
        FileNotFoundError: If the workbook path does not exist.
        ValueError: If the workbook format or profile is unsupported.
        ImportError: If ``openpyxl`` is not installed.
    """
    workbook_path = _resolve_workbook_path(path)
    openpyxl = _require_openpyxl()

    normalized_key = _normalize_family_key(family_key)
    profile_name = _resolve_profile(normalized_key, workbook_path)
    family_name = family_name or DEFAULT_FAMILY_NAMES.get(profile_name, family_key)

    workbook = openpyxl.load_workbook(
        workbook_path,
        data_only=True,
        read_only=False,
    )
    try:
        spec = PROFILES[profile_name]
        source_id = f"{normalized_key}:{workbook_path.stem}"

        route_sheet = _select_sheet(workbook, spec["process_flow"]["sheets"], "process_flow")
        pfmea_sheet = _select_sheet(workbook, spec["pfmea"]["sheets"], "pfmea")
        control_sheet = _select_sheet(
            workbook, spec["control_plan"]["sheets"], "control_plan"
        )

        route_items = _parse_process_route(
            route_sheet,
            spec["process_flow"],
            workbook_path,
            source_id,
        )
        pfmea_items = _parse_pfmea_patterns(
            pfmea_sheet,
            spec["pfmea"],
            workbook_path,
            source_id,
        )
        control_items = _parse_control_plan_patterns(
            control_sheet,
            spec["control_plan"],
            workbook_path,
            source_id,
        )
        special_items = _collect_special_characteristics(
            route_items,
            pfmea_items,
            control_items,
        )
    finally:
        workbook.close()

    route_items = _assign_ids(
        _dedupe_items(
            route_items,
            (
                "process_step_id",
                "operation",
                "process_function",
                "special_characteristics",
            ),
        ),
        "template_id",
        f"{normalized_key}-route",
    )
    pfmea_items = _assign_ids(
        _dedupe_items(
            pfmea_items,
            (
                "process_step_id",
                "process_function",
                "failure_mode",
                "failure_effect",
                "failure_cause",
                "prevention_control",
                "detection_control",
                "classification",
                "rating_hint",
            ),
        ),
        "pattern_id",
        f"{normalized_key}-pfmea",
    )
    control_items = _assign_ids(
        _dedupe_items(
            control_items,
            (
                "process_step_id",
                "operation",
                "characteristic",
                "classification",
                "specification",
                "measurement_method",
                "sample_size",
                "sample_frequency",
                "control_method",
                "reaction_plan",
            ),
        ),
        "pattern_id",
        f"{normalized_key}-control-plan",
    )
    special_items = _assign_ids(
        _dedupe_items(
            special_items,
            ("kind", "value", "related_process", "related_characteristic"),
        ),
        "pattern_id",
        f"{normalized_key}-special",
    )

    return {
        "schema_version": SCHEMA_VERSION,
        "family_key": normalized_key,
        "family_name": family_name,
        "process_route_templates": route_items,
        "pfmea_patterns": pfmea_items,
        "control_plan_patterns": control_items,
        "special_characteristic_patterns": special_items,
        "source_refs": _dedupe_source_refs(
            [
                _workbook_source_ref(workbook_path, source_id),
                *_collect_refs(route_items),
                *_collect_refs(pfmea_items),
                *_collect_refs(control_items),
                *_collect_refs(special_items),
            ]
        ),
    }


def load_part_family_knowledge(
    family_key_or_path: str | Path | None = None,
) -> dict[str, Any]:
    """Load part-family knowledge by family key, JSON/XLSX path, or all defaults.

    - ``None`` returns a mapping of all indexed default family payloads.
    - A known family key (for example ``stamping`` or ``门铰链``) loads that family.
    - A JSON file loads and normalizes the payload.
    - A workbook file or seed directory is converted in memory.

    Args:
        family_key_or_path: One of ``None``, a family key string, a JSON path,
            or a workbook / seed-directory path.

    Returns:
        A single knowledge payload dict, or a dict of family_key -> payload
        when ``family_key_or_path`` is ``None``.

    Raises:
        FileNotFoundError: If the given path does not exist.
        KeyError: If the family key is unknown and no seed directory is found.
        ValueError: If the file type is unsupported.
    """
    if family_key_or_path is None:
        index = load_part_family_index()
        return {
            entry["family_key"]: _load_index_entry(entry, _index_path().parent)
            for entry in index.get("families", [])
            if isinstance(entry, dict) and entry.get("family_key")
        }

    candidate = Path(family_key_or_path)
    if candidate.exists():
        if candidate.is_dir():
            index_path = candidate / "index.json"
            if index_path.exists():
                index = load_part_family_index(index_path)
                return {
                    entry["family_key"]: _load_index_entry(entry, index_path.parent)
                    for entry in index.get("families", [])
                    if isinstance(entry, dict) and entry.get("family_key")
                }
            family_key = _infer_family_key_from_path(candidate)
            return convert_seed_workbook_to_family_knowledge(
                candidate,
                family_key,
                DEFAULT_FAMILY_NAMES.get(family_key),
            )
        if candidate.suffix.lower() == ".json":
            return _normalize_family_payload(_read_json(candidate), candidate)
        if candidate.suffix.lower() in {".xlsx", ".xlsm"}:
            family_key = _infer_family_key_from_path(candidate)
            return convert_seed_workbook_to_family_knowledge(
                candidate,
                family_key,
                DEFAULT_FAMILY_NAMES.get(family_key),
            )
        raise ValueError(f"不支持的零件族知识文件类型: {candidate}")

    text = str(family_key_or_path)
    if _looks_like_path(text):
        raise FileNotFoundError(f"零件族知识路径不存在: {candidate}")

    key = _normalize_family_key(text)
    index = load_part_family_index()
    for entry in index.get("families", []):
        if not isinstance(entry, dict):
            continue
        if _normalize_family_key(str(entry.get("family_key", ""))) == key:
            return _load_index_entry(entry, _index_path().parent)

    seed_dir = _default_seed_directories().get(key)
    if seed_dir and seed_dir.exists():
        return convert_seed_workbook_to_family_knowledge(
            seed_dir,
            key,
            DEFAULT_FAMILY_NAMES.get(key),
        )

    raise KeyError(f"未知零件族知识库: {family_key_or_path}")


def load_part_family_index(path: str | Path | None = None) -> dict[str, Any]:
    """Load the part-family index or build one from default configuration.

    Args:
        path: Optional explicit path to an index JSON file or directory
            containing ``index.json``. Defaults to the standard index path.

    Returns:
        An index dict with ``schema_version`` and a ``families`` list.

    Raises:
        FileNotFoundError: If an explicit path is given but does not exist.
        ValueError: If the JSON content is malformed.
    """
    index_path = Path(path) if path is not None else _index_path()
    if index_path.is_dir():
        index_path = index_path / "index.json"

    if index_path.exists():
        payload = _read_json(index_path)
        if not isinstance(payload, dict):
            raise ValueError(f"零件族索引必须是 JSON 对象: {index_path}")
        families = payload.get("families", [])
        if not isinstance(families, list):
            raise ValueError(f"零件族索引 families 必须是列表: {index_path}")
        normalized = []
        for entry in families:
            if isinstance(entry, dict):
                normalized.append(_normalize_index_entry(entry))
        return {
            "schema_version": str(
                payload.get("schema_version", INDEX_SCHEMA_VERSION)
            ),
            "families": normalized,
        }

    if path is not None:
        raise FileNotFoundError(f"零件族索引不存在: {index_path}")

    return _build_default_index()


def _load_index_entry(entry: dict[str, Any], base_dir: Path) -> dict[str, Any]:
    """Load a single index entry by resolving its knowledge_path or seed directory."""
    knowledge_path = str(entry.get("knowledge_path", "")).strip()
    if knowledge_path:
        candidate = Path(knowledge_path)
        if not candidate.is_absolute():
            candidate = base_dir / candidate
        if candidate.exists():
            return _normalize_family_payload(_read_json(candidate), candidate)

    key = _normalize_family_key(str(entry.get("family_key", "")))
    seed_dir_text = str(entry.get("default_seed_dir", "")).strip()
    seed_dir = Path(seed_dir_text) if seed_dir_text else _default_seed_directories().get(key)
    if seed_dir and seed_dir.exists():
        return convert_seed_workbook_to_family_knowledge(
            seed_dir,
            key,
            str(entry.get("family_name") or DEFAULT_FAMILY_NAMES.get(key) or key),
        )

    raise FileNotFoundError(
        f"零件族知识库不存在: {entry.get('family_key')} "
        f"(knowledge_path={knowledge_path!r}, default_seed_dir={seed_dir})"
    )


def _normalize_family_payload(payload: Any, path: Path) -> dict[str, Any]:
    """Normalize a raw JSON family payload into a consistent structure."""
    if not isinstance(payload, dict):
        raise ValueError(f"零件族知识库必须是 JSON 对象: {path}")

    family_key = _normalize_family_key(str(payload.get("family_key", path.stem)))
    family_name = str(payload.get("family_name") or DEFAULT_FAMILY_NAMES.get(family_key) or family_key)
    normalized: dict[str, Any] = {
        "schema_version": str(payload.get("schema_version", SCHEMA_VERSION)),
        "family_key": family_key,
        "family_name": family_name,
        "process_route_templates": _assign_ids(
            _dedupe_items(
                _ensure_list(payload.get("process_route_templates")),
                (
                    "process_step_id",
                    "operation",
                    "process_function",
                    "special_characteristics",
                ),
            ),
            "template_id",
            f"{family_key}-route",
        ),
        "pfmea_patterns": _assign_ids(
            _dedupe_items(
                _ensure_list(payload.get("pfmea_patterns")),
                (
                    "process_step_id",
                    "process_function",
                    "failure_mode",
                    "failure_effect",
                    "failure_cause",
                    "prevention_control",
                    "detection_control",
                    "classification",
                    "rating_hint",
                ),
            ),
            "pattern_id",
            f"{family_key}-pfmea",
        ),
        "control_plan_patterns": _assign_ids(
            _dedupe_items(
                _ensure_list(payload.get("control_plan_patterns")),
                (
                    "process_step_id",
                    "operation",
                    "characteristic",
                    "classification",
                    "specification",
                    "measurement_method",
                    "sample_size",
                    "sample_frequency",
                    "control_method",
                    "reaction_plan",
                ),
            ),
            "pattern_id",
            f"{family_key}-control-plan",
        ),
        "special_characteristic_patterns": _assign_ids(
            _dedupe_items(
                _ensure_list(payload.get("special_characteristic_patterns")),
                ("kind", "value", "related_process", "related_characteristic"),
            ),
            "pattern_id",
            f"{family_key}-special",
        ),
        "source_refs": _dedupe_source_refs(_ensure_list(payload.get("source_refs"))),
    }
    if not normalized["source_refs"]:
        normalized["source_refs"] = _dedupe_source_refs(
            [
                *_collect_refs(normalized["process_route_templates"]),
                *_collect_refs(normalized["pfmea_patterns"]),
                *_collect_refs(normalized["control_plan_patterns"]),
                *_collect_refs(normalized["special_characteristic_patterns"]),
            ]
        )
    process_inference = _load_process_inference_payload(family_key, path)
    if process_inference:
        normalized["process_inference"] = process_inference
        normalized["process_route_templates"] = _assign_ids(
            _dedupe_items(
                normalized["process_route_templates"]
                + _process_inference_route_templates(process_inference, family_key),
                (
                    "process_step_id",
                    "operation",
                    "process_function",
                    "special_characteristics",
                ),
            ),
            "template_id",
            f"{family_key}-route",
        )
        normalized["source_refs"] = _dedupe_source_refs(
            normalized["source_refs"] + _ensure_list(process_inference.get("source_refs"))
        )
    return normalized


def _load_process_inference_payload(
    family_key: str,
    family_payload_path: Path,
) -> dict[str, Any]:
    """Load the process-inference JSON for a family, if one exists under the config tree."""
    config_dir = _config_dir()
    if not family_payload_path.resolve().is_relative_to(config_dir.resolve()):
        return {}
    candidates: list[Path] = []
    if family_payload_path.suffix.lower() == ".json":
        root = family_payload_path.parent.parent / "process_inference"
        candidates.append(root / f"{family_key}.json")
    candidates.append(_process_inference_dir() / f"{family_key}.json")
    for candidate in candidates:
        if not candidate.exists():
            continue
        payload = _read_json(candidate)
        if not isinstance(payload, dict):
            raise ValueError(f"工艺推断知识库必须是 JSON 对象: {candidate}")
        if _normalize_family_key(str(payload.get("family_key", family_key))) != family_key:
            raise ValueError(f"工艺推断知识库 family_key 不匹配: {candidate}")
        normalized = dict(payload)
        normalized["source_refs"] = _dedupe_source_refs(
            _ensure_list(payload.get("source_refs"))
            + [
                {
                    "source_id": f"process_inference:{family_key}",
                    "path": str(candidate),
                    "knowledge_type": "process_inference",
                }
            ]
        )
        return normalized
    return {}


def _process_inference_route_templates(
    process_inference: dict[str, Any],
    family_key: str,
) -> list[dict[str, Any]]:
    """Convert process-inference variants into route-template dicts."""
    variants = _variant_items(process_inference.get("process_flow_variants"))
    variants.extend(_variant_items(process_inference.get("process_flow_variants_extended")))
    source_refs = _ensure_list(process_inference.get("source_refs"))
    templates: list[dict[str, Any]] = []
    for variant_id, variant in variants:
        sequence = _ensure_list(variant.get("process_sequence"))
        if not sequence:
            continue
        templates.append(
            {
                "template_id": f"{family_key}-process-inference-{variant_id}",
                "variant_id": variant_id,
                "variant_description": str(variant.get("description", "")).strip(),
                "steps": [
                    _process_inference_step(step, index)
                    for index, step in enumerate(sequence, start=1)
                    if isinstance(step, dict)
                ],
                "knowledge_source_refs": [f"process_inference:{family_key}:{variant_id}"],
                "source_refs": source_refs,
                "source": "process_inference",
                "confidence": 0.8,
                "needs_confirmation": True,
            }
        )
    return [template for template in templates if template["steps"]]


def _variant_items(value: Any) -> list[tuple[str, dict[str, Any]]]:
    """Extract (variant_id, variant_dict) pairs from a raw value."""
    if not isinstance(value, dict):
        return []
    items: list[tuple[str, dict[str, Any]]] = []
    for key, row in value.items():
        if isinstance(row, dict):
            items.append((str(key).strip() or f"variant_{len(items) + 1}", row))
    return items


def _process_inference_step(step: dict[str, Any], index: int) -> dict[str, Any]:
    """Build a single route-template step dict from a process-inference entry."""
    seq = _integer(str(step.get("seq", ""))) or index * 10
    operation = _clean(step.get("operation"))
    key_check = _clean(step.get("key_check"))
    return {
        "process_step_id": f"P{seq:02d}" if seq < 100 else f"P{seq}",
        "sequence": seq,
        "operation": operation,
        "process_function": key_check or operation,
        "process_purpose": key_check or operation,
        "equipment": _clean(step.get("equipment")),
        "process_characteristics": [key_check] if key_check else [],
        "reasoning": "根据工艺推断知识库的零件族工艺变体补充 PFD 工序",
        "knowledge_source_refs": [],
        "confidence": 0.8,
        "needs_confirmation": True,
        "preserve_operation_name": True,
    }


def _parse_process_route(
    worksheet: Any,
    spec: dict[str, Any],
    workbook_path: Path,
    source_id: str,
) -> list[dict[str, Any]]:
    """Parse process-flow rows from an openpyxl worksheet using a column spec."""
    merges = _merged_ranges(worksheet)
    items: list[dict[str, Any]] = []
    generated_index = 1
    for row_index in _row_range(worksheet, spec):
        step_text = _value(worksheet, row_index, spec["step"], merges)
        if not _is_data_text(step_text):
            continue
        step_id = _step_id(step_text, generated_index)
        sequence = _sequence_from_id(step_id, generated_index * 10)
        operation = _first_nonempty(
            _operation_from_step(step_text),
            _combine(worksheet, row_index, spec.get("operation", []), merges),
            f"Legacy process step {sequence}",
        )
        process_function = (
            _combine(worksheet, row_index, spec.get("function", []), merges)
            or operation
        )
        specials = _unique(
            _values(worksheet, row_index, spec.get("special", []), merges)
        )
        items.append(
            {
                "process_step_id": step_id,
                "sequence": sequence,
                "operation": operation,
                "process_function": process_function,
                "special_characteristics": specials,
                "source_refs": [
                    _row_source_ref(workbook_path, worksheet.title, row_index, source_id)
                ],
            }
        )
        generated_index += 1
    return items


def _parse_pfmea_patterns(
    worksheet: Any,
    spec: dict[str, Any],
    workbook_path: Path,
    source_id: str,
) -> list[dict[str, Any]]:
    """Parse PFMEA pattern rows from an openpyxl worksheet using a column spec."""
    merges = _merged_ranges(worksheet)
    state: dict[str, str] = {}
    items: list[dict[str, Any]] = []
    generated_index = 1
    for row_index in _row_range(worksheet, spec):
        _update_state(worksheet, row_index, spec.get("inherit", []), merges, state)
        failure_mode = _value(worksheet, row_index, spec["failure_mode"], merges)
        if not _is_data_text(failure_mode):
            continue
        step_text = state.get(spec["step"], "") or _value(
            worksheet, row_index, spec["step"], merges
        )
        step_id = _step_id(step_text, generated_index)
        severity = _integer(_value(worksheet, row_index, spec["severity"], merges))
        occurrence = _integer(_value(worksheet, row_index, spec["occurrence"], merges))
        detection = _integer(_value(worksheet, row_index, spec["detection"], merges))
        classification = _value(worksheet, row_index, spec["classification"], merges)
        action_priority = _value(
            worksheet, row_index, spec.get("action_priority", ""), merges
        )
        if not action_priority and (classification or severity >= 9):
            action_priority = "H" if severity >= 9 else "M"
        items.append(
            {
                "process_step_id": step_id,
                "process_function": _combine_with_state(
                    worksheet,
                    row_index,
                    spec.get("function", []),
                    merges,
                    state,
                ),
                "failure_mode": failure_mode,
                "failure_effect": _value(worksheet, row_index, spec["effect"], merges),
                "failure_cause": _value(worksheet, row_index, spec["cause"], merges),
                "prevention_control": _value(
                    worksheet, row_index, spec["prevention"], merges
                ),
                "detection_control": _value(
                    worksheet,
                    row_index,
                    spec["detection_control"],
                    merges,
                ),
                "classification": classification,
                "rating_hint": {
                    "severity": severity,
                    "occurrence": occurrence,
                    "detection": detection,
                    "action_priority": action_priority,
                },
                "source_refs": [
                    _row_source_ref(workbook_path, worksheet.title, row_index, source_id)
                ],
            }
        )
        generated_index += 1
    return items


def _parse_control_plan_patterns(
    worksheet: Any,
    spec: dict[str, Any],
    workbook_path: Path,
    source_id: str,
) -> list[dict[str, Any]]:
    """Parse control-plan pattern rows from an openpyxl worksheet using a column spec."""
    merges = _merged_ranges(worksheet)
    state: dict[str, str] = {}
    items: list[dict[str, Any]] = []
    last_explicit_by_step: dict[str, dict[str, Any]] = {}
    generated_index = 1
    for row_index in _row_range(worksheet, spec):
        _update_state(worksheet, row_index, spec.get("inherit", []), merges, state)
        characteristic = _combine(
            worksheet, row_index, spec.get("characteristic", []), merges
        )
        has_core_fields = _has_any(
            worksheet,
            row_index,
            spec,
            ["specification", "measurement", "control_method"],
            merges,
        )
        if not _is_data_text(characteristic) and not has_core_fields:
            continue

        step_text = state.get(spec["step"], "") or _value(
            worksheet, row_index, spec["step"], merges
        )
        step_id = _step_id(step_text, generated_index)
        specification = _value(worksheet, row_index, spec["specification"], merges)
        inherited_from = None
        previous_explicit = last_explicit_by_step.get(step_id)
        if not characteristic and not specification and previous_explicit:
            characteristic = previous_explicit["characteristic"]
            specification = previous_explicit["specification"]
            inherited_from = previous_explicit

        item: dict[str, Any] = {
            "process_step_id": step_id,
            "operation": _combine(worksheet, row_index, spec.get("operation", []), merges)
            or state.get(spec.get("operation", ""), ""),
            "characteristic": characteristic,
            "classification": _value(
                worksheet, row_index, spec["classification"], merges
            ),
            "specification": specification,
            "measurement_method": _value(
                worksheet, row_index, spec["measurement"], merges
            ),
            "sample_size": _value(worksheet, row_index, spec["sample_size"], merges),
            "sample_frequency": _value(
                worksheet, row_index, spec["sample_frequency"], merges
            ),
            "control_method": _value(
                worksheet, row_index, spec["control_method"], merges
            ),
            "reaction_plan": _value(worksheet, row_index, spec["reaction_plan"], merges),
            "source_refs": [
                _row_source_ref(workbook_path, worksheet.title, row_index, source_id)
            ],
        }
        if inherited_from:
            item["continuation_of_source_row"] = inherited_from["source_refs"][0].get("row")
        items.append(item)
        if characteristic and specification and not inherited_from:
            last_explicit_by_step[step_id] = item
        generated_index += 1
    return items


def _collect_special_characteristics(
    route_items: list[dict[str, Any]],
    pfmea_items: list[dict[str, Any]],
    control_items: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Aggregate special-characteristic patterns from all three knowledge sources."""
    items: list[dict[str, Any]] = []
    for route in route_items:
        for value in _ensure_string_list(route.get("special_characteristics")):
            items.append(
                {
                    "kind": "process_flow_special_characteristic",
                    "value": value,
                    "related_process": route.get("operation", ""),
                    "related_characteristic": "",
                    "source_refs": _ensure_list(route.get("source_refs")),
                }
            )
    for pattern in pfmea_items:
        classification = str(pattern.get("classification", "")).strip()
        if classification:
            items.append(
                {
                    "kind": "pfmea_classification",
                    "value": classification,
                    "related_process": pattern.get("process_function", ""),
                    "related_characteristic": pattern.get("failure_mode", ""),
                    "source_refs": _ensure_list(pattern.get("source_refs")),
                }
            )
    for control in control_items:
        classification = str(control.get("classification", "")).strip()
        if classification:
            items.append(
                {
                    "kind": "control_plan_classification",
                    "value": classification,
                    "related_process": control.get("operation", ""),
                    "related_characteristic": control.get("characteristic", ""),
                    "source_refs": _ensure_list(control.get("source_refs")),
                }
            )
    return items


def _require_openpyxl() -> Any:
    """Import and return the ``openpyxl`` module, raising a user-friendly error if missing."""
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "解析 XLSX 零件族知识库需要安装 openpyxl；请安装 openpyxl 后重试。"
        ) from exc
    return openpyxl


def _resolve_workbook_path(path: str | Path) -> Path:
    """Resolve a path to a single .xlsx/.xlsm workbook, expanding directories."""
    candidate = Path(path)
    if candidate.is_dir():
        workbooks = sorted(
            item
            for item in candidate.iterdir()
            if item.is_file()
            and item.suffix.lower() in {".xlsx", ".xlsm"}
            and not item.name.startswith("~$")
        )
        if not workbooks:
            raise FileNotFoundError(f"种子目录中没有 XLSX/XLSM 文件: {candidate}")
        return workbooks[0]
    if not candidate.exists():
        raise FileNotFoundError(f"种子 workbook 不存在: {candidate}")
    if candidate.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"种子 workbook 必须是 .xlsx 或 .xlsm: {candidate}")
    return candidate


def _resolve_profile(family_key: str, workbook_path: Path) -> str:
    """Resolve the profile name for a family key, inferring from path if needed."""
    key = _normalize_family_key(family_key)
    if key in PROFILES:
        return key
    inferred = _infer_family_key_from_path(workbook_path)
    if inferred in PROFILES:
        return inferred
    raise ValueError(f"不支持的零件族: {family_key}")


def _infer_family_key_from_path(path: Path) -> str:
    """Guess a family key from a path by matching against FAMILY_ALIASES."""
    text = f"{path.parent.name} {path.stem} {path.name}"
    for alias, key in FAMILY_ALIASES.items():
        if alias.lower() in text.lower() or alias in text:
            return key
    normalized = _normalize_family_key(path.stem)
    return normalized if normalized in PROFILES else path.stem


def _normalize_family_key(value: str) -> str:
    """Normalise a family key string by compacting whitespace and resolving aliases."""
    compact = _compact(value).lower().replace("-", "_")
    return FAMILY_ALIASES.get(compact, FAMILY_ALIASES.get(value, compact))


def _select_sheet(workbook: Any, names: list[str], kind: str) -> Any:
    """Select the best-matching worksheet from a workbook by name or kind."""
    compact_to_name = {_compact(name): name for name in workbook.sheetnames}
    for candidate in names:
        compact_candidate = _compact(candidate)
        for compact_name, sheet_name in compact_to_name.items():
            if compact_candidate and compact_candidate in compact_name:
                return workbook[sheet_name]
    for sheet_name in workbook.sheetnames:
        compact = _compact(sheet_name).lower()
        if kind == "process_flow" and ("过程流程" in compact or "flow" in compact):
            return workbook[sheet_name]
        if kind == "pfmea" and "pfmea" in compact:
            return workbook[sheet_name]
        if kind == "control_plan" and (
            "控制计划" in compact or "controlplan" in compact
        ):
            return workbook[sheet_name]
    raise ValueError(
        f"找不到 {kind} 工作表；可用工作表: {', '.join(workbook.sheetnames)}"
    )


def _row_range(worksheet: Any, spec: dict[str, Any]) -> range:
    """Return the integer row range for parsing based on the column spec."""
    start = int(spec.get("start", 1))
    end = int(spec.get("end") or worksheet.max_row or start)
    return range(start, min(end, int(worksheet.max_row or end)) + 1)


def _merged_ranges(worksheet: Any) -> list[tuple[int, int, int, int, Any]]:
    """Extract merged-cell ranges and their values from an openpyxl worksheet."""
    ranges: list[tuple[int, int, int, int, Any]] = []
    merged_cells = getattr(worksheet, "merged_cells", None)
    for merged in getattr(merged_cells, "ranges", ()):
        min_col, min_row, max_col, max_row = merged.bounds
        value = worksheet.cell(min_row, min_col).value
        ranges.append((min_row, min_col, max_row, max_col, value))
    return ranges


def _value(
    worksheet: Any,
    row_index: int,
    column: str,
    merges: list[tuple[int, int, int, int, Any]],
) -> str:
    """Read a single cell value, falling back to merged-cell values."""
    if not column:
        return ""
    col_index = _column_index(column)
    value = worksheet.cell(row_index, col_index).value
    if value is None:
        for min_row, min_col, max_row, max_col, merged_value in merges:
            if min_row <= row_index <= max_row and min_col <= col_index <= max_col:
                value = merged_value
                break
    return _clean(value)


def _values(
    worksheet: Any,
    row_index: int,
    columns: str | list[str],
    merges: list[tuple[int, int, int, int, Any]],
) -> list[str]:
    """Read multiple cell values for a given row."""
    if isinstance(columns, str):
        columns = [columns]
    return [
        value
        for column in columns
        if (value := _value(worksheet, row_index, column, merges))
    ]


def _combine(
    worksheet: Any,
    row_index: int,
    columns: str | list[str],
    merges: list[tuple[int, int, int, int, Any]],
) -> str:
    """Combine multiple column values into a single semicolon-separated string."""
    return "；".join(_unique(_values(worksheet, row_index, columns, merges)))


def _combine_with_state(
    worksheet: Any,
    row_index: int,
    columns: str | list[str],
    merges: list[tuple[int, int, int, int, Any]],
    state: dict[str, str],
) -> str:
    """Combine column values with inherited state fallback."""
    if isinstance(columns, str):
        columns = [columns]
    values = []
    for column in columns:
        values.append(_value(worksheet, row_index, column, merges) or state.get(column, ""))
    return "；".join(_unique([value for value in values if value]))


def _update_state(
    worksheet: Any,
    row_index: int,
    columns: list[str],
    merges: list[tuple[int, int, int, int, Any]],
    state: dict[str, str],
) -> None:
    """Update an inherited-state dict from the specified columns of the current row."""
    for column in columns:
        value = _value(worksheet, row_index, column, merges)
        if value:
            state[column] = value


def _has_any(
    worksheet: Any,
    row_index: int,
    spec: dict[str, Any],
    keys: list[str],
    merges: list[tuple[int, int, int, int, Any]],
) -> bool:
    """Check whether any of the specified spec keys have a non-empty value in the current row."""
    return any(_value(worksheet, row_index, spec[key], merges) for key in keys if spec.get(key))


def _assign_ids(
    items: list[dict[str, Any]],
    field: str,
    prefix: str,
) -> list[dict[str, Any]]:
    """Assign sequential IDs to items that do not already have one."""
    result: list[dict[str, Any]] = []
    for index, item in enumerate(items, start=1):
        copy = dict(item)
        copy[field] = str(copy.get(field) or f"{prefix}-{index:03d}")
        result.append(copy)
    return result


def _dedupe_items(
    items: list[dict[str, Any]],
    keys: tuple[str, ...],
) -> list[dict[str, Any]]:
    """Deduplicate items by a tuple of key fields, merging source_refs on duplicates."""
    result: list[dict[str, Any]] = []
    seen: dict[tuple[Any, ...], dict[str, Any]] = {}
    for item in items:
        if not isinstance(item, dict):
            continue
        key = tuple(_hashable_value(item.get(name)) for name in keys)
        if key in seen:
            existing = seen[key]
            existing["source_refs"] = _dedupe_source_refs(
                [
                    *_ensure_list(existing.get("source_refs")),
                    *_ensure_list(item.get("source_refs")),
                ]
            )
            continue
        copy = dict(item)
        copy["source_refs"] = _dedupe_source_refs(_ensure_list(copy.get("source_refs")))
        seen[key] = copy
        result.append(copy)
    return result


def _dedupe_source_refs(refs: list[Any]) -> list[dict[str, Any]]:
    """Deduplicate a list of source-ref dicts, removing empty fields."""
    result: list[dict[str, Any]] = []
    seen: set[tuple[tuple[str, str], ...]] = set()
    for ref in refs:
        if not isinstance(ref, dict):
            continue
        cleaned = {str(k): v for k, v in ref.items() if v not in ("", None, [])}
        key = tuple(sorted((name, str(value)) for name, value in cleaned.items()))
        if key in seen:
            continue
        seen.add(key)
        result.append(cleaned)
    return result


def _collect_refs(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Collect all source_ref entries from a list of item dicts."""
    refs: list[dict[str, Any]] = []
    for item in items:
        refs.extend(_ensure_list(item.get("source_refs")))
    return refs


def _workbook_source_ref(path: Path, source_id: str) -> dict[str, Any]:
    """Build a top-level source-ref dict for a workbook."""
    return {
        "source_id": source_id,
        "path": str(path),
        "workbook": path.name,
    }


def _row_source_ref(
    path: Path,
    sheet: str,
    row: int,
    source_id: str,
) -> dict[str, Any]:
    """Build a row-level source-ref dict for a workbook cell."""
    return {
        "source_id": source_id,
        "path": str(path),
        "workbook": path.name,
        "sheet": sheet,
        "row": row,
    }


def _normalize_index_entry(entry: dict[str, Any]) -> dict[str, Any]:
    """Normalise a single index entry by resolving aliases and stripping empty fields."""
    key = _normalize_family_key(str(entry.get("family_key", "")))
    normalized = {
        "family_key": key,
        "family_name": str(entry.get("family_name") or DEFAULT_FAMILY_NAMES.get(key) or key),
    }
    for field in ("knowledge_path", "default_seed_dir", "default_seed_workbook"):
        value = str(entry.get(field, "")).strip()
        if value:
            normalized[field] = value
    return normalized


def _build_default_index() -> dict[str, Any]:
    """Build the default part-family index from built-in profiles and seed directories."""
    families: list[dict[str, Any]] = []
    for key in ("stamping", "limiter", "door_hinge"):
        entry: dict[str, Any] = {
            "family_key": key,
            "family_name": DEFAULT_FAMILY_NAMES[key],
        }
        knowledge_path = _config_dir() / f"{key}.json"
        if knowledge_path.exists():
            entry["knowledge_path"] = knowledge_path.name
        seed_dir = _default_seed_directories()[key]
        entry["default_seed_dir"] = str(seed_dir)
        try:
            seed_workbook = _resolve_workbook_path(seed_dir)
        except (FileNotFoundError, ValueError):
            seed_workbook = None
        if seed_workbook is not None:
            entry["default_seed_workbook"] = seed_workbook.name
        families.append(entry)
    return {"schema_version": INDEX_SCHEMA_VERSION, "families": families}


def _read_json(path: Path) -> Any:
    """Read and parse a UTF-8 JSON file."""
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def _looks_like_path(value: str) -> bool:
    """Check whether a string looks like a filesystem path."""
    return bool(re.search(r"[\\/]", value) or Path(value).suffix)


def _hashable_value(value: Any) -> Any:
    """Convert a value into a hashable form for deduplication."""
    if isinstance(value, dict):
        return tuple(sorted((str(k), _hashable_value(v)) for k, v in value.items()))
    if isinstance(value, list):
        return tuple(_hashable_value(item) for item in value)
    if isinstance(value, tuple):
        return tuple(_hashable_value(item) for item in value)
    return _clean(value).lower() if isinstance(value, str) else value


def _step_id(text: str, fallback_index: int) -> str:
    """Extract a step ID (e.g. P10) from text, generating a fallback if none is found."""
    match = re.search(r"(?<!\d)(\d{1,3})(?!\d)", text or "")
    if match:
        number = int(match.group(1))
        width = 2 if number < 100 else 3
        return f"P{number:0{width}d}"
    return f"P{fallback_index * 10:02d}"


def _sequence_from_id(step_id: str, default: int) -> int:
    """Extract the numeric sequence from a step ID string."""
    match = re.search(r"\d+", step_id)
    return int(match.group(0)) if match else default


def _operation_from_step(text: str) -> str:
    """Strip leading numbering from a step text to extract the operation name."""
    return _clean(re.sub(r"^\s*\d+\s*[-_.、:]?\s*", "", text or ""))


def _is_data_text(text: str) -> bool:
    """Check whether a text value looks like meaningful data rather than a header/noise."""
    value = _clean(text)
    if not value:
        return False
    upper = value.upper()
    noise = (
        "END",
        "PFMEA",
        "控制计划",
        "过程流程",
        "编号",
        "项目",
        "日期",
        "表单",
        "OPERATION NO",
        "FAILURE MODE",
    )
    return not any(item in upper for item in noise)


def _integer(value: str) -> int:
    """Parse the first integer out of a string, returning 0 if none is found."""
    match = re.search(r"-?\d+", value or "")
    return int(match.group(0)) if match else 0


def _column_index(column: str) -> int:
    """Convert an Excel column letter (e.g. 'AB') to a 1-based integer index."""
    index = 0
    for char in column.upper():
        if not ("A" <= char <= "Z"):
            continue
        index = index * 26 + ord(char) - ord("A") + 1
    return index


def _ensure_list(value: Any) -> list[Any]:
    """Return the value if it is a list, otherwise an empty list."""
    return value if isinstance(value, list) else []


def _ensure_string_list(value: Any) -> list[str]:
    """Coerce a value to a list of cleaned non-empty strings."""
    if isinstance(value, list):
        return [_clean(item) for item in value if _clean(item)]
    if value:
        return [_clean(value)]
    return []


def _unique(values: list[str]) -> list[str]:
    """Deduplicate a list of strings while preserving order."""
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = _clean(value)
        if cleaned and cleaned not in seen:
            seen.add(cleaned)
            result.append(cleaned)
    return result


def _first_nonempty(*values: str) -> str:
    """Return the first non-empty string from a sequence of values."""
    for value in values:
        if value:
            return value
    return ""


def _clean(value: Any) -> str:
    """Normalise whitespace and convert a value to a stripped string."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\r", "\n")
    text = re.sub(r"[ \t　]+", " ", text)
    text = re.sub(r"\n+", "\n", text)
    return text.strip()


def _compact(value: str) -> str:
    """Remove all whitespace from a string."""
    return re.sub(r"\s+", "", value or "")
