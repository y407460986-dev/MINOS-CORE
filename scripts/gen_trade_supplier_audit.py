import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# 样式定义
title_font = Font(name="微软雅黑", size=16, bold=True)
header_font = Font(name="微软雅黑", size=11, bold=True)
normal_font = Font(name="微软雅黑", size=10)
small_font = Font(name="微软雅黑", size=9)
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font_white = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
light_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_title(ws, row, title, col_span):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = title_font
    cell.alignment = center_align
    ws.row_dimensions[row].height = 40


def write_section_header(ws, row, title, col_span):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = header_font_white
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.fill = header_fill
    for c in range(1, col_span + 1):
        ws.cell(row=row, column=c).border = thin_border
        if c > 1:
            ws.cell(row=row, column=c).fill = header_fill
    ws.row_dimensions[row].height = 28


def write_header_row(ws, row, headers):
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = Font(name="微软雅黑", size=10, bold=True)
        cell.alignment = center_align
        cell.border = thin_border
        cell.fill = light_fill


def write_data_row(ws, row, data, fill=None):
    for i, d in enumerate(data, 1):
        cell = ws.cell(row=row, column=i, value=d)
        cell.font = normal_font
        cell.alignment = left_align if i > 1 else center_align
        cell.border = thin_border
        if fill:
            cell.fill = fill


def write_kv_rows(ws, start_row, items, col_span):
    """写入键值对行: [(label, value), ...]"""
    r = start_row
    for label, value in items:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        cell_l = ws.cell(row=r, column=1, value=label)
        cell_l.font = Font(name="微软雅黑", size=10, bold=True)
        cell_l.alignment = left_align
        cell_l.border = thin_border
        cell_l.fill = yellow_fill
        ws.cell(row=r, column=2).border = thin_border
        ws.cell(row=r, column=2).fill = yellow_fill

        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=col_span)
        cell_v = ws.cell(row=r, column=3, value=value)
        cell_v.font = normal_font
        cell_v.alignment = left_align
        cell_v.border = thin_border
        for c in range(4, col_span + 1):
            ws.cell(row=r, column=c).border = thin_border
        r += 1
    return r


# ===================== Sheet 1: 审核表主表 =====================
ws1 = wb.active
ws1.title = "贸易商准入审核表"
set_col_widths(ws1, [6, 22, 18, 18, 18, 18, 18])

COL = 7
r = 1

# 标题
write_title(ws1, r, "贸易商准入审核表（汽车零部件）", COL)
r += 1
ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COL)
c = ws1.cell(row=r, column=1, value="文件编号：QA-SRM-TSA-001    版本：V1.0    生效日期：2026-06-25")
c.font = small_font
c.alignment = center_align
r += 2

# 一、基本信息
write_section_header(ws1, r, "一、基本信息", COL)
r += 1
kv_items = [
    ("贸易商名称", ""),
    ("统一社会信用代码", ""),
    ("注册地址", ""),
    ("实际经营地址", ""),
    ("法定代表人", ""),
    ("注册资本（万元）", ""),
    ("成立日期", ""),
    ("经营范围", ""),
    ("联系人", ""),
    ("联系电话", ""),
    ("电子邮箱", ""),
    ("主营产品类别", ""),
    ("代理品牌/厂家", ""),
    ("申请准入日期", ""),
    ("申请人", ""),
]
r = write_kv_rows(ws1, r, kv_items, COL)
r += 1

# 二、资质文件审核
write_section_header(ws1, r, "二、资质文件审核", COL)
r += 1
write_header_row(ws1, r, ["序号", "审核项目", "是否齐全", "有效期至", "备注", "", ""])
r += 1
qual_items = [
    ("1", "营业执照（副本）", "□ 是  □ 否", "", "", "", ""),
    ("2", "一般纳税人资格证明", "□ 是  □ 否", "", "", "", ""),
    ("3", "开户许可证", "□ 是  □ 否", "", "", "", ""),
    ("4", "法人身份证复印件", "□ 是  □ 否", "", "", "", ""),
    ("5", "授权委托书（如适用）", "□ 是  □ 否", "", "", "", ""),
    ("6", "品牌授权书/代理证书", "□ 是  □ 否", "", "", "", ""),
    ("7", "产品检验报告/合格证", "□ 是  □ 否", "", "", "", ""),
    ("8", "进出口资质（如适用）", "□ 是  □ 否", "", "", "", ""),
    ("9", "危险品经营许可（如适用）", "□ 是  □ 否", "", "", "", ""),
    ("10", "IATF 16949 体系证书（上游厂家）", "□ 是  □ 否", "", "", "", ""),
    ("11", "ISO 9001 质量管理体系证书", "□ 是  □ 否", "", "", "", ""),
    ("12", "ISO 14001 环境管理体系证书", "□ 是  □ 否", "", "", "", ""),
    ("13", "产品材料安全数据表 MSDS（如适用）", "□ 是  □ 否", "", "", "", ""),
]
for item in qual_items:
    write_data_row(ws1, r, item)
    r += 1

ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
c = ws1.cell(row=r, column=1, value="资质审核结论")
c.font = Font(name="微软雅黑", size=10, bold=True)
c.alignment = left_align
c.fill = yellow_fill
c.border = thin_border
ws1.cell(row=r, column=2).border = thin_border
ws1.cell(row=r, column=2).fill = yellow_fill
ws1.merge_cells(start_row=r, start_column=3, end_row=r, end_column=COL)
c2 = ws1.cell(row=r, column=3, value="□ 通过    □ 有条件通过    □ 不通过")
c2.font = normal_font
c2.alignment = left_align
c2.border = thin_border
for cc in range(4, COL + 1):
    ws1.cell(row=r, column=cc).border = thin_border
r += 2

# 三~六：评分表
score_sections = [
    (
        "三、质量保证能力",
        ["序号", "评估项目", "评估内容", "评分(1-5)", "备注", "", ""],
        [
            ("1", "质量管理体系", "是否建立完善的质量管理体系，是否有专职质量管理人员"),
            ("2", "进货检验能力", "是否具备来料检验能力，检验设备/手段是否齐全"),
            ("3", "不合格品管控", "是否有不合格品处理流程及追溯机制"),
            ("4", "产品追溯能力", "是否能提供完整的产品批次追溯记录"),
            ("5", "质量问题响应", "质量投诉响应时效及处理机制"),
            ("6", "包装防护", "产品包装是否满足汽车零部件防护要求（防锈、防尘、防碰撞）"),
            ("7", "储存条件", "仓储环境是否满足产品储存要求（温湿度、防潮等）"),
            ("8", "变更管理", "供应商/产品变更时是否主动通知并提供变更资料"),
        ],
        "质量保证能力评分：_____ / 40 分    评分标准：1=不满足  2=基本不满足  3=基本满足  4=满足  5=优秀",
    ),
    (
        "四、供货能力评估",
        ["序号", "评估项目", "评估内容", "评分(1-5)", "备注", "", ""],
        [
            ("1", "库存管理", "常用品种是否有安全库存，库存周转率"),
            ("2", "交货准时率", "承诺交期内准时交付的能力"),
            ("3", "紧急订单响应", "紧急/临时订单的响应速度及满足率"),
            ("4", "物流配送能力", "配送方式、配送时效、物流覆盖范围"),
            ("5", "供货稳定性", "批次间质量一致性，供货连续性"),
            ("6", "多品种管理", "SKU管理能力，品种齐全度"),
            ("7", "最小起订量", "MOQ是否灵活，是否支持小批量供货"),
        ],
        "供货能力评分：_____ / 35 分",
    ),
    (
        "五、财务与商务能力",
        ["序号", "评估项目", "评估内容", "评分(1-5)", "备注", "", ""],
        [
            ("1", "注册资金规模", "注册资本是否与经营规模匹配"),
            ("2", "营业收入", "近三年年均营业收入水平"),
            ("3", "账期承受能力", "是否接受约定账期（如月结30/60/90天）"),
            ("4", "价格竞争力", "同等品质下价格水平及价格稳定性"),
            ("5", "发票开具", "开票及时性、税率准确性"),
        ],
        "财务与商务评分：_____ / 25 分",
    ),
    (
        "六、服务能力评估",
        ["序号", "评估项目", "评估内容", "评分(1-5)", "备注", "", ""],
        [
            ("1", "技术支持", "是否能提供产品技术资料、选型建议"),
            ("2", "售后服务", "退换货政策及执行效率"),
            ("3", "沟通响应", "日常沟通效率，问题反馈时效"),
            ("4", "信息共享", "市场行情、价格波动等信息的主动共享"),
            ("5", "合规配合", "配合审核、提供资料的意愿及效率"),
        ],
        "服务能力评分：_____ / 25 分",
    ),
]

for sec_title, sec_headers, sec_items, sec_summary in score_sections:
    write_section_header(ws1, r, sec_title, COL)
    r += 1
    write_header_row(ws1, r, sec_headers)
    r += 1
    for item in sec_items:
        write_data_row(ws1, r, list(item) + [""] * (COL - len(item)))
        r += 1
    ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COL)
    c = ws1.cell(row=r, column=1, value=sec_summary)
    c.font = Font(name="微软雅黑", size=10, bold=True, color="4472C4")
    c.alignment = left_align
    r += 2

# 七、汽车行业专项
write_section_header(ws1, r, "七、汽车行业专项要求", COL)
r += 1
write_header_row(ws1, r, ["序号", "检查项", "是否满足", "说明", "", "", ""])
r += 1
auto_items = [
    ("1", "所代理产品是否通过相关汽车行业认证（如CCC、ECE、DOT等）"),
    ("2", "是否了解PPAP（生产件批准程序）要求"),
    ("3", "是否能配合执行8D/PDCA等质量问题处理流程"),
    ("4", "是否具备IMDS/CAMDS材料数据提交能力（如适用）"),
    ("5", "是否了解REACH/ELV等环保法规要求"),
    ("6", "是否能提供符合汽车行业要求的产品标识及追溯标签"),
    ("7", "是否同意签署质量协议及保密协议"),
    ("8", "是否接受年度绩效考核及现场审核"),
]
for num, item in auto_items:
    write_data_row(ws1, r, [num, item, "□ 是  □ 否", "", "", "", ""])
    r += 1
r += 1

# 八、风险评估
write_section_header(ws1, r, "八、风险评估", COL)
r += 1
write_header_row(ws1, r, ["风险类别", "风险描述", "风险等级", "应对措施", "", "", ""])
r += 1
risk_items = [
    ("经营风险", "贸易商经营稳定性、是否存在诉讼/失信记录"),
    ("供应风险", "上游厂家依赖度、单一货源风险"),
    ("质量风险", "产品假冒伪劣风险、以次充好风险"),
    ("合规风险", "资质造假、环保/安全合规风险"),
    ("廉洁风险", "商业贿赂、利益输送风险"),
]
for cat, desc in risk_items:
    write_data_row(ws1, r, [cat, desc, "□ 高  □ 中  □ 低", "", "", "", ""])
    r += 1
r += 1

# 九、综合评分汇总
write_section_header(ws1, r, "九、综合评分汇总", COL)
r += 1
write_header_row(ws1, r, ["评估维度", "满分", "得分", "权重", "加权得分", "", ""])
r += 1
summary_items = [
    ("质量保证能力", "40", "", "30%", ""),
    ("供货能力", "35", "", "25%", ""),
    ("财务与商务能力", "25", "", "20%", ""),
    ("服务能力", "25", "", "15%", ""),
    ("汽车行业专项", "—", "", "10%", ""),
]
for item in summary_items:
    write_data_row(ws1, r, list(item) + ["", ""])
    r += 1
write_data_row(ws1, r, ["合计", "125", "", "100%", "", "", ""], fill=yellow_fill)
r += 1

# 准入等级表
write_header_row(ws1, r, ["等级", "综合得分", "准入结论", "", "", "", ""])
r += 1
grade_items = [
    ("A级（优秀）", "≥90分", "直接准入，优先采购"),
    ("B级（合格）", "75-89分", "准入，正常采购"),
    ("C级（有条件）", "60-74分", "有条件准入，限期整改，降级采购"),
    ("D级（不合格）", "<60分", "不准入"),
]
for g in grade_items:
    write_data_row(ws1, r, list(g) + ["", "", "", ""])
    r += 1
r += 1

# 十、审核结论
write_section_header(ws1, r, "十、审核结论", COL)
r += 1
conclusion_items = [
    ("准入等级", "□ A级    □ B级    □ C级    □ D级"),
    ("审核结论", "□ 准入    □ 有条件准入    □ 不准入"),
    ("有条件准入整改要求", ""),
    ("整改期限", ""),
    ("限制采购范围（如适用）", ""),
]
r = write_kv_rows(ws1, r, conclusion_items, COL)
r += 1

# 十一、审批签字
write_section_header(ws1, r, "十一、审批签字", COL)
r += 1
write_header_row(ws1, r, ["审批环节", "签字", "日期", "意见", "", "", ""])
r += 1
approval_items = [
    "采购部初审",
    "质量部审核",
    "技术部审核（如需）",
    "财务部审核",
    "分管领导审批",
    "总经理审批（A级/C级）",
]
for item in approval_items:
    write_data_row(ws1, r, [item, "", "", "", "", "", ""])
    r += 1
r += 1

# 十二、附件清单
write_section_header(ws1, r, "十二、附件清单", COL)
r += 1
attach_items = [
    "营业执照副本复印件",
    "品牌授权书/代理证书复印件",
    "质量管理体系证书复印件",
    "产品检验报告",
    "企业征信报告",
    "近三年财务报表（资产负债表、利润表）",
    "主要客户清单",
    "质量协议（签署版）",
    "保密协议（签署版）",
    "廉洁承诺书（签署版）",
    "现场审核报告（如适用）",
    "其他：_______________",
]
for _i, item in enumerate(attach_items, 1):
    write_data_row(ws1, r, ["", f"□  {item}", "", "", "", "", ""])
    r += 1
r += 1

# 附录
write_section_header(ws1, r, "附录：贸易商年度复评要求", COL)
r += 1
write_header_row(ws1, r, ["复评项目", "说明", "", "", "", "", ""])
r += 1
review_items = [
    ("复评周期", "每年一次，准入满12个月前30天启动"),
    ("复评依据", "年度供货业绩、质量表现、交期达成率、配合度"),
    ("绩效考核指标", "交货准时率≥95%、来料合格率≥99%、投诉响应≤24h"),
    ("降级/淘汰规则", "连续两次C级降为淘汰；发生重大质量事故直接淘汰"),
    ("资质更新", "复评时需更新所有即将到期的资质文件"),
]
for proj, desc in review_items:
    write_data_row(ws1, r, [proj, desc, "", "", "", "", ""])
    r += 1

# 尾注
r += 1
ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COL)
c = ws1.cell(row=r, column=1, value="本审核表由采购部归档管理，保存期限不少于贸易商合作终止后3年。")
c.font = small_font
c.alignment = left_align

# 打印设置
ws1.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
ws1.page_setup.fitToWidth = 1
ws1.page_setup.fitToHeight = 0
ws1.page_setup.orientation = "portrait"

output_path = r"D:\MINOS-CORE\docs\贸易商准入审核表_汽车零部件.xlsx"
wb.save(output_path)
print(f"已生成: {output_path}")
