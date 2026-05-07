from __future__ import annotations

import asyncio
import tempfile
import unittest
from pathlib import Path

from httpx import ASGITransport, AsyncClient

from minos_desktop.api.app import create_app
from minos_desktop.services import DesktopRunService

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None  # type: ignore[assignment]


class DesktopApiV114Tests(unittest.TestCase):
    def test_closure_tasks_api_serializes_and_resolves_fake_run(self) -> None:
        asyncio.run(self._closure_tasks_api_serializes_and_resolves_fake_run())

    def test_knowledge_family_api_lists_default_families(self) -> None:
        asyncio.run(self._knowledge_family_api_lists_default_families())

    @unittest.skipIf(Workbook is None, "openpyxl is required for seed import API tests")
    def test_knowledge_import_seeds_api_converts_three_families(self) -> None:
        asyncio.run(self._knowledge_import_seeds_api_converts_three_families())

    async def _closure_tasks_api_serializes_and_resolves_fake_run(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            db_path = Path(tmpdir) / "desktop.sqlite3"
            output_dir = Path(tmpdir) / "out"
            app = create_app(db_path=db_path, output_dir=output_dir)
            run = DesktopRunService(db_path).save_result(None, _fake_result(), output_dir)
            transport = ASGITransport(app=app)
            async with AsyncClient(transport=transport, base_url="http://test") as client:
                drawing = await client.get(f"/api/runs/{run.id}/drawing-analysis")
                part_family = await client.get(f"/api/runs/{run.id}/part-family")
                updated_family = await client.patch(
                    f"/api/runs/{run.id}/part-family",
                    json={"part_family": "stamping", "family_name": "冲压件", "reason": "人工确认"},
                )
                tasks = await client.get(f"/api/runs/{run.id}/closure-tasks")
                confirm = await client.patch(
                    f"/api/runs/{run.id}/process-flow-confirmations",
                    json={"step_ids": ["P10"], "reason": "PFD 已人工确认"},
                )

        self.assertEqual(drawing.status_code, 200)
        self.assertEqual(drawing.json()["materials"][0]["item_id"], "MAT-1")
        self.assertEqual(drawing.json()["run_id"], run.id)
        self.assertEqual(part_family.status_code, 200)
        self.assertEqual(part_family.json()["part_family"], "未设置")
        self.assertEqual(updated_family.status_code, 200)
        self.assertEqual(updated_family.json()["part_family"], "stamping")
        self.assertEqual(updated_family.json()["context"]["family_name"], "冲压件")
        self.assertEqual(tasks.status_code, 200)
        body = tasks.json()
        self.assertEqual(body["run_id"], run.id)
        self.assertTrue(
            any(task["task_type"] == "process_flow_confirmation" for task in body["tasks"])
        )
        self.assertEqual(confirm.status_code, 200)
        updated = confirm.json()["current_result"]["process_flow"]["steps"][0]
        self.assertFalse(updated["needs_confirmation"])
        self.assertEqual(updated["review_status"], "已确认")

    async def _knowledge_family_api_lists_default_families(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            db_path = Path(tmpdir) / "desktop.sqlite3"
            output_dir = Path(tmpdir) / "out"
            app = create_app(db_path=db_path, output_dir=output_dir)
            transport = ASGITransport(app=app)
            async with AsyncClient(transport=transport, base_url="http://test") as client:
                response = await client.get("/api/knowledge/families")

        self.assertEqual(response.status_code, 200)
        keys = {item["family_key"] for item in response.json()["families"]}
        self.assertGreaterEqual(keys, {"stamping", "limiter", "door_hinge"})

    async def _knowledge_import_seeds_api_converts_three_families(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            seed_dirs = {
                "stamping": str(_build_stamping_seed(root / "stamping")),
                "limiter": str(_build_limiter_seed(root / "limiter")),
                "door_hinge": str(_build_door_hinge_seed(root / "door_hinge")),
            }
            db_path = root / "desktop.sqlite3"
            output_dir = root / "out"
            app = create_app(db_path=db_path, output_dir=output_dir)
            transport = ASGITransport(app=app)
            async with AsyncClient(transport=transport, base_url="http://test") as client:
                response = await client.post(
                    "/api/knowledge/import-seeds",
                    json={"seed_dirs": seed_dirs, "output_dir": str(root / "knowledge")},
                )

            self.assertEqual(response.status_code, 200)
            body = response.json()
            self.assertEqual(body["status"], "ok")
            families = {item["family_key"]: item for item in body["families"]}
            self.assertGreaterEqual(set(families), {"stamping", "limiter", "door_hinge"})
            for item in families.values():
                self.assertGreaterEqual(item["process_route_template_count"], 1)
                self.assertGreaterEqual(item["pfmea_pattern_count"], 1)
                self.assertGreaterEqual(item["control_plan_pattern_count"], 1)
                self.assertTrue(Path(item["output_path"]).exists())


def _fake_result() -> dict:
    return {
        "document_type": "three_table_generation_result",
        "status": "needs_review",
        "drawing_analysis": {
            "materials": [
                {
                    "item_id": "MAT-1",
                    "material": "SPCC",
                    "confidence": 0.6,
                    "needs_confirmation": True,
                    "evidence_refs": ["EV-1"],
                }
            ],
            "evidence": [{"evidence_id": "EV-1", "content": "material callout"}],
        },
        "process_flow": {
            "steps": [
                {
                    "process_step_id": "P10",
                    "operation": "来料检验",
                    "needs_confirmation": True,
                    "evidence_refs": ["EV-1"],
                }
            ]
        },
        "pfmea": {"rows": []},
        "control_plan": {"rows": []},
        "metadata": {},
        "manual_review": {"required": True, "items": []},
        "exports": {},
        "manifest": {"run_id": "fake-run", "path": "manifest.json"},
    }


def _build_stamping_seed(root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    path = root / "1004258-冲压件.xlsx"
    workbook = Workbook()
    process = workbook.active
    process.title = "7 过程流程图"
    pfmea = workbook.create_sheet("8 PFMEA")
    control = workbook.create_sheet("9 控制计划")
    process["A9"] = "10 冲压"
    process["I9"] = "SC"
    pfmea["A16"] = "10 冲压"
    pfmea["B16"] = "冲压成形"
    pfmea["C16"] = "毛刺"
    pfmea["D16"] = "装配干涉"
    pfmea["E16"] = 8
    pfmea["F16"] = "SC"
    pfmea["G16"] = "模具磨损"
    pfmea["H16"] = 4
    pfmea["I16"] = "模具点检"
    pfmea["J16"] = "巡检"
    pfmea["K16"] = 5
    control["A13"] = "10"
    control["B13"] = "冲压"
    control["F13"] = "毛刺"
    control["I13"] = "SC"
    control["J13"] = "无毛刺"
    control["K13"] = "目视"
    control["L13"] = "5件"
    control["M13"] = "每批"
    control["N13"] = "巡检"
    control["O13"] = "隔离并通知质量"
    workbook.save(path)
    workbook.close()
    return root


def _build_limiter_seed(root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    path = root / "XWQ01000100-限位器.xlsx"
    workbook = Workbook()
    process = workbook.active
    process.title = "4过程流程图"
    pfmea = workbook.create_sheet("6PFMEA")
    control = workbook.create_sheet("7控制计划-过程")
    process["A10"] = "10"
    process["B10"] = "铆接"
    process["G10"] = "限位器装配"
    process["H10"] = "KPC"
    pfmea["B10"] = "10"
    pfmea["D10"] = "铆接"
    pfmea["I10"] = "铆接高度不良"
    pfmea["G10"] = "功能失效"
    pfmea["H10"] = 9
    pfmea["P10"] = "SC"
    pfmea["J10"] = "设备压力波动"
    pfmea["L10"] = 4
    pfmea["K10"] = "设备点检"
    pfmea["M10"] = "高度检测"
    pfmea["N10"] = 4
    pfmea["O10"] = "H"
    control["A12"] = "10"
    control["B12"] = "铆接"
    control["I12"] = "铆接高度"
    control["O12"] = "SC"
    control["Q12"] = "按图纸"
    control["U12"] = "高度规"
    control["Y12"] = "5件"
    control["Z12"] = "每小时"
    control["AB12"] = "SPC"
    control["AE12"] = "停机隔离"
    workbook.save(path)
    workbook.close()
    return root


def _build_door_hinge_seed(root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    path = root / "JLH01000300-门铰链.xlsx"
    workbook = Workbook()
    process = workbook.active
    process.title = "过程流程图"
    pfmea = workbook.create_sheet("PFMEA")
    control = workbook.create_sheet("控制计划")
    process["A8"] = "10 进料检验"
    process["F8"] = "产品检验"
    process["H8"] = "KPC 尺寸"
    pfmea["A16"] = "10 进料检验"
    pfmea["B16"] = "尺寸超差"
    pfmea["D16"] = "装配干涉"
    pfmea["F16"] = 8
    pfmea["G16"] = "SC"
    pfmea["H16"] = "来料尺寸波动"
    pfmea["I16"] = 4
    pfmea["J16"] = "供应商过程审核"
    pfmea["K16"] = "入库抽检"
    pfmea["L16"] = 5
    control["A16"] = "10"
    control["B16"] = "进料检验"
    control["E16"] = "安装孔尺寸"
    control["G16"] = "SC"
    control["H16"] = "按图纸公差"
    control["I16"] = "卡尺"
    control["J16"] = "5件"
    control["K16"] = "每批"
    control["L16"] = "抽检"
    control["N16"] = "隔离并通知QE"
    workbook.save(path)
    workbook.close()
    return root


if __name__ == "__main__":
    unittest.main()
