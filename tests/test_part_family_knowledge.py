"""Tests for XLSX seed conversion into part-family knowledge."""

from __future__ import annotations

import builtins
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from minos_core.part_family_knowledge import (
    convert_seed_workbook_to_family_knowledge,
    load_part_family_index,
    load_part_family_knowledge,
)

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None  # type: ignore[assignment]


def _write_json(payload: dict) -> Path:
    fd, path = tempfile.mkstemp(suffix=".json", text=True)
    os.close(fd)
    target = Path(path)
    target.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return target


@unittest.skipIf(Workbook is None, "openpyxl is required for workbook conversion tests")
class ConvertSeedWorkbookTests(unittest.TestCase):
    def test_convert_seed_workbook_to_family_knowledge(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "JLH01000300-门铰链.xlsx"
            self._build_workbook(workbook_path)

            knowledge = convert_seed_workbook_to_family_knowledge(
                workbook_path,
                "door_hinge",
            )

        self.assertEqual(knowledge["family_key"], "door_hinge")
        self.assertEqual(knowledge["family_name"], "门铰链")
        self.assertIn("process_route_templates", knowledge)
        self.assertIn("pfmea_patterns", knowledge)
        self.assertIn("control_plan_patterns", knowledge)
        self.assertIn("special_characteristic_patterns", knowledge)
        self.assertIn("source_refs", knowledge)
        self.assertEqual(len(knowledge["process_route_templates"]), 2)
        self.assertEqual(len(knowledge["pfmea_patterns"]), 1)
        self.assertEqual(len(knowledge["control_plan_patterns"]), 1)

        pfmea = knowledge["pfmea_patterns"][0]
        self.assertEqual(pfmea["failure_mode"], "尺寸超差")
        self.assertEqual(pfmea["rating_hint"]["severity"], 8)
        self.assertEqual(len(pfmea["source_refs"]), 2)
        self.assertEqual({ref["row"] for ref in pfmea["source_refs"]}, {16, 17})

        control = knowledge["control_plan_patterns"][0]
        self.assertEqual(control["characteristic"], "安装孔尺寸")
        self.assertEqual(len(control["source_refs"]), 2)
        self.assertTrue(
            any(
                ref.get("sheet") == "PFMEA" and ref.get("row") == 16
                for ref in knowledge["source_refs"]
            )
        )
        self.assertTrue(
            any(
                item["value"] == "SC"
                for item in knowledge["special_characteristic_patterns"]
            )
        )

    def test_convert_accepts_seed_directory(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "JLH01000300-门铰链.xlsx"
            self._build_workbook(workbook_path)

            knowledge = convert_seed_workbook_to_family_knowledge(root, "门铰链")

        self.assertEqual(knowledge["family_key"], "door_hinge")
        self.assertEqual(knowledge["family_name"], "门铰链")
        self.assertGreaterEqual(len(knowledge["source_refs"]), 1)

    def _build_workbook(self, path: Path) -> None:
        workbook = Workbook()
        process = workbook.active
        process.title = "过程流程图"
        pfmea = workbook.create_sheet("PFMEA")
        control_plan = workbook.create_sheet("控制计划")

        process["A8"] = "10 进料检验"
        process["F8"] = "产品检验"
        process["H8"] = "KPC 尺寸"
        process["K8"] = "KCC 扭矩"
        process["A9"] = "20 装配"
        process["F9"] = "铰链装配"

        for row in (16, 17):
            pfmea[f"A{row}"] = "10 进料检验"
            pfmea[f"B{row}"] = "尺寸超差"
            pfmea[f"D{row}"] = "装配干涉"
            pfmea[f"F{row}"] = 8
            pfmea[f"G{row}"] = "SC"
            pfmea[f"H{row}"] = "来料尺寸波动"
            pfmea[f"I{row}"] = 4
            pfmea[f"J{row}"] = "供应商过程审核"
            pfmea[f"K{row}"] = "入库抽检"
            pfmea[f"L{row}"] = 5

        for row in (16, 17):
            control_plan[f"A{row}"] = "10"
            control_plan[f"B{row}"] = "进料检验"
            control_plan[f"E{row}"] = "安装孔尺寸"
            control_plan[f"G{row}"] = "SC"
            control_plan[f"H{row}"] = "按图纸公差"
            control_plan[f"I{row}"] = "卡尺"
            control_plan[f"J{row}"] = "5件"
            control_plan[f"K{row}"] = "每批"
            control_plan[f"L{row}"] = "抽检"
            control_plan[f"N{row}"] = "隔离并通知QE"

        workbook.save(path)
        workbook.close()


class LoadPartFamilyKnowledgeTests(unittest.TestCase):
    def test_load_json_path_deduplicates_and_preserves_source_refs(self):
        payload = {
            "family_key": "stamping",
            "family_name": "冲压件",
            "process_route_templates": [],
            "pfmea_patterns": [
                {
                    "process_step_id": "P10",
                    "process_function": "冲压",
                    "failure_mode": "毛刺",
                    "failure_effect": "装配干涉",
                    "failure_cause": "模具磨损",
                    "prevention_control": "点检",
                    "detection_control": "巡检",
                    "classification": "",
                    "rating_hint": {"severity": 7, "occurrence": 4, "detection": 5},
                    "source_refs": [{"workbook": "seed.xlsx", "sheet": "PFMEA", "row": 16}],
                },
                {
                    "process_step_id": "P10",
                    "process_function": "冲压",
                    "failure_mode": "毛刺",
                    "failure_effect": "装配干涉",
                    "failure_cause": "模具磨损",
                    "prevention_control": "点检",
                    "detection_control": "巡检",
                    "classification": "",
                    "rating_hint": {"severity": 7, "occurrence": 4, "detection": 5},
                    "source_refs": [{"workbook": "seed.xlsx", "sheet": "PFMEA", "row": 17}],
                },
            ],
            "control_plan_patterns": [],
            "special_characteristic_patterns": [],
            "source_refs": [],
        }
        path = _write_json(payload)
        try:
            knowledge = load_part_family_knowledge(path)
        finally:
            path.unlink()

        self.assertEqual(knowledge["family_key"], "stamping")
        self.assertEqual(len(knowledge["pfmea_patterns"]), 1)
        refs = knowledge["pfmea_patterns"][0]["source_refs"]
        self.assertEqual({ref["row"] for ref in refs}, {16, 17})
        self.assertEqual({ref["row"] for ref in knowledge["source_refs"]}, {16, 17})

    def test_load_part_family_index_from_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            index_path = root / "index.json"
            index_path.write_text(
                json.dumps(
                    {
                        "schema_version": "custom",
                        "families": [
                            {
                                "family_key": "门铰链",
                                "family_name": "门铰链",
                                "knowledge_path": "door_hinge.json",
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            index = load_part_family_index(index_path)

        self.assertEqual(index["schema_version"], "custom")
        self.assertEqual(index["families"][0]["family_key"], "door_hinge")
        self.assertEqual(index["families"][0]["knowledge_path"], "door_hinge.json")

    def test_default_index_lists_three_seed_families(self):
        index = load_part_family_index()
        keys = {entry["family_key"] for entry in index["families"]}
        self.assertGreaterEqual(keys, {"stamping", "limiter", "door_hinge"})

    def test_default_family_loads_process_inference_context(self):
        knowledge = load_part_family_knowledge("door_hinge")

        self.assertIn("process_inference", knowledge)
        process_inference = knowledge["process_inference"]
        self.assertEqual(process_inference["family_key"], "door_hinge")
        self.assertIn("process_flow_variants", process_inference)
        self.assertTrue(
            any(
                template.get("source") == "process_inference"
                for template in knowledge["process_route_templates"]
            )
        )

    def test_openpyxl_missing_raises_clear_exception(self):
        fd, raw_path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        workbook_path = Path(raw_path)

        real_import = builtins.__import__

        def blocked_import(name, *args, **kwargs):
            if name == "openpyxl" or name.startswith("openpyxl."):
                raise ImportError("blocked openpyxl")
            return real_import(name, *args, **kwargs)

        try:
            with mock.patch("builtins.__import__", side_effect=blocked_import):
                with self.assertRaisesRegex(ImportError, "openpyxl"):
                    convert_seed_workbook_to_family_knowledge(
                        workbook_path,
                        "stamping",
                    )
        finally:
            workbook_path.unlink()


if __name__ == "__main__":
    unittest.main()
