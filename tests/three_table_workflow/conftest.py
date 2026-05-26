import json
import tempfile
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image


class SharedTestHelpers:
    """Shared test helpers for three table workflow tests."""

    @staticmethod
    def write_png(path: Path) -> Path:
        Image.new("RGB", (100, 80), color="white").save(path)
        return path

    @staticmethod
    def analysis_payload():
        refs = ["drawing_page_001"]
        return {
            "drawing_identity": {
                "part_name": "Hinge Bracket",
                "drawing_number": "HB-001",
                "revision": "A",
                "evidence_refs": refs,
            },
            "materials": [
                {"name": "SPCC", "confidence": 0.92, "evidence_refs": refs}
            ],
            "surface_treatments": [
                {"name": "Zinc plated", "confidence": 0.8, "evidence_refs": refs}
            ],
            "dimensions": [
                {"name": "hole diameter", "value": "8.0 +/-0.1", "confidence": 0.88, "evidence_refs": refs}
            ],
            "tolerances": [],
            "special_characteristics": [
                {
                    "id": "SC-001",
                    "name": "hole position",
                    "classification": "SC",
                    "specification": "+/-0.1",
                    "confidence": 0.7,
                    "evidence_refs": refs,
                }
            ],
            "technical_requirements": [
                {"text": "No burrs", "confidence": 0.9, "evidence_refs": refs}
            ],
            "inspection_requirements": [
                {"text": "Check critical dimensions", "confidence": 0.85, "evidence_refs": refs}
            ],
            "manufacturing_clues": [
                {
                    "process_step_id": "P10",
                    "operation": "Blanking",
                    "process_function": "Create blank profile",
                    "equipment": "Press",
                    "confidence": 0.82,
                    "evidence_refs": refs,
                },
                {
                    "process_step_id": "P20",
                    "operation": "Piercing",
                    "process_function": "Create mounting holes",
                    "equipment": "Piercing die",
                    "special_characteristics": ["SC-001"],
                    "confidence": 0.78,
                    "evidence_refs": refs,
                },
                {
                    "process_step_id": "P30",
                    "operation": "Surface treatment",
                    "process_function": "Apply corrosion protection",
                    "confidence": 0.76,
                    "evidence_refs": refs,
                },
            ],
            "uncertainties": [
                {"text": "Exact process sequence requires engineer review", "confidence": 0.6, "evidence_refs": refs}
            ],
        }

    @staticmethod
    def assert_xlsx(test_case, path: Path):
        test_case.assertTrue(Path(path).exists())
        workbook = load_workbook(path, data_only=False, read_only=True)
        try:
            test_case.assertTrue(workbook.sheetnames)
            for sheet_name in workbook.sheetnames:
                workbook[sheet_name]
        finally:
            workbook.close()

    @staticmethod
    def xlsx_text(path: Path) -> str:
        with ZipFile(path) as archive:
            return "\n".join(
                archive.read(name).decode("utf-8")
                for name in archive.namelist()
                if name.endswith(".xml")
            )

    @staticmethod
    def sheet_names(path: Path) -> list[str]:
        workbook = load_workbook(path, data_only=False, read_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()


def canonical_review_key(value):
    text = str(value or "").strip()
    if text.startswith("failure_mode:"):
        return "pfmea:" + text.split(":", 1)[1]
    if text.startswith("control_item:"):
        return "control_plan:" + text.split(":", 1)[1]
    if text.startswith("process_step:"):
        return "pfd:" + text.split(":", 1)[1]
    if text.startswith("PFMEA-"):
        return "pfmea:" + text
    if text.startswith("CP-"):
        return "control_plan:" + text
    return text


def rows_by_step(rows):
    by_step = {}
    for row in rows:
        by_step.setdefault(row["process_step_id"], []).append(row)
    return by_step


PFMEA_CONTROL_LOGIC_FIELDS = (
    "cause_process_parameter",
    "prevention_control_target",
    "prevention_control_timing",
    "detection_control_target",
    "detection_control_timing",
    "control_logic_review",
)


def assert_pfmea_ai_risk_row(test_case, row):
    required_fields = {
        "process_item",
        "process_step",
        "process_work_element",
        "structure_process_item",
        "structure_process_step",
        "structure_work_element",
        "function_process_item",
        "function_process_step",
        "function_work_element",
        "failure_effect",
        "failure_mode",
        "failure_cause",
        "severity",
        "occurrence",
        "detection",
        "action_priority",
        "S",
        "O",
        "D",
        "AP",
        "prevention_control",
        "detection_control",
        *PFMEA_CONTROL_LOGIC_FIELDS,
        "preventive_action",
        "detection_action",
        "optimization_action",
        "recommended_action",
        "responsible_person",
        "target_date",
        "status",
        "ai_reasoning",
        "confidence",
        "ai_inferred",
        "needs_confirmation",
        "evidence_refs",
    }
    test_case.assertTrue(required_fields.issubset(row))
    test_case.assertIn(row["action_priority"], {"H", "M", "L"})
    test_case.assertTrue(row["prevention_control"])
    test_case.assertTrue(row["detection_control"])
    for field in PFMEA_CONTROL_LOGIC_FIELDS:
        test_case.assertTrue(row[field], field)
    test_case.assertTrue(row["optimization_action"])
    test_case.assertTrue(row["recommended_action"])
    test_case.assertTrue(row["ai_reasoning"])
    test_case.assertTrue(row["ai_inferred"])
    test_case.assertNotIn("rpn", row)
    test_case.assertNotIn("after_rpn", row)
    test_case.assertIsInstance(row["needs_confirmation"], bool)
    test_case.assertIsInstance(row["evidence_refs"], list)
    test_case.assertGreaterEqual(row["confidence"], 0.0)
    test_case.assertLessEqual(row["confidence"], 0.9)
